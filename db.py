"""
Umumiy ma'lumotlar bazasi qatlami.
Bot (main.py) va Web App backend (app.py) ikkalasi ham shu modulni ishlatadi,
shunda ular bitta attendance.db faylida bir xil mantiq bilan ishlaydi.
"""
import os
import math
import sqlite3
from datetime import datetime
from zoneinfo import ZoneInfo

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

TASHKENT_TZ = ZoneInfo("Asia/Tashkent")
DB_PATH = os.getenv("DB_PATH", "attendance.db")


def get_now():
    return datetime.now(TASHKENT_TZ)


def get_db():
    conn = sqlite3.connect(DB_PATH, timeout=10)
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def calculate_distance(lat1, lon1, lat2, lon2):
    R = 6371000
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    delta_phi = math.radians(lat2 - lat1)
    delta_lambda = math.radians(lon2 - lon1)
    a = math.sin(delta_phi / 2) ** 2 + math.cos(phi1) * math.cos(phi2) * math.sin(delta_lambda / 2) ** 2
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
    return R * c


def init_db(admin_ids):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS users (
                user_id INTEGER PRIMARY KEY,
                full_name TEXT,
                is_approved INTEGER DEFAULT 0,
                monthly_salary REAL DEFAULT 0.0,
                norm_days INTEGER DEFAULT 26
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS attendance (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                date TEXT,
                check_in_time TEXT,
                check_out_time TEXT,
                lateness_minutes INTEGER DEFAULT 0,
                lateness_reason TEXT DEFAULT '',
                break_start TEXT DEFAULT NULL,
                break_end TEXT DEFAULT NULL,
                break_minutes INTEGER DEFAULT 0,
                work_hours REAL DEFAULT 0,
                FOREIGN KEY(user_id) REFERENCES users(user_id)
            )
        ''')
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS advances (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER,
                amount REAL DEFAULT 0.0,
                date TEXT,
                FOREIGN KEY(user_id) REFERENCES users(user_id)
            )
        ''')
        for table, col in [
            ("users", "is_approved INTEGER DEFAULT 0"),
            ("users", "monthly_salary REAL DEFAULT 0.0"),
            ("users", "norm_days INTEGER DEFAULT 26"),
            ("attendance", "lateness_reason TEXT DEFAULT ''"),
            ("attendance", "break_start TEXT DEFAULT NULL"),
            ("attendance", "break_end TEXT DEFAULT NULL"),
            ("attendance", "break_minutes INTEGER DEFAULT 0"),
        ]:
            try:
                cursor.execute(f"ALTER TABLE {table} ADD COLUMN {col}")
            except sqlite3.OperationalError:
                pass
        try:
            cursor.execute(
                "CREATE UNIQUE INDEX IF NOT EXISTS idx_attendance_user_date "
                "ON attendance(user_id, date)"
            )
        except sqlite3.IntegrityError:
            pass
        for admin_id in admin_ids:
            cursor.execute(
                "INSERT OR IGNORE INTO users (user_id, full_name, is_approved, monthly_salary, norm_days) "
                "VALUES (?, ?, 1, 0.0, 26)", (admin_id, "Admin")
            )
            cursor.execute("UPDATE users SET is_approved = 1 WHERE user_id = ?", (admin_id,))
        conn.commit()


# ---------------- USERS ----------------
def get_user(user_id):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT user_id, full_name, is_approved, monthly_salary, norm_days FROM users WHERE user_id = ?",
            (user_id,)
        )
        return cursor.fetchone()


def is_user_approved(user_id, admin_ids):
    if user_id in admin_ids:
        return True
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT is_approved FROM users WHERE user_id = ?", (user_id,))
        row = cursor.fetchone()
        return bool(row and row[0] == 1)


def insert_new_user(user_id, full_name, is_approved):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "INSERT OR IGNORE INTO users (user_id, full_name, is_approved, monthly_salary, norm_days) "
            "VALUES (?, ?, ?, 0, 26)", (user_id, full_name, is_approved)
        )
        conn.commit()


def approve_user(user_id):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("UPDATE users SET is_approved = 1 WHERE user_id = ?", (user_id,))
        conn.commit()


def reject_user(user_id):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM users WHERE user_id = ?", (user_id,))
        conn.commit()


def get_all_users():
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT user_id, full_name, is_approved, monthly_salary, norm_days FROM users ORDER BY full_name COLLATE NOCASE")
        return cursor.fetchall()


def get_approved_users():
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT user_id, full_name FROM users WHERE is_approved = 1")
        return cursor.fetchall()


def rename_user(user_id, new_name):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("UPDATE users SET full_name = ? WHERE user_id = ?", (new_name, user_id))
        conn.commit()


def upsert_user(user_id, name):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "INSERT INTO users (user_id, full_name, is_approved, monthly_salary, norm_days) VALUES (?, ?, 1, 0.0, 26) "
            "ON CONFLICT(user_id) DO UPDATE SET full_name = excluded.full_name, is_approved = 1",
            (user_id, name)
        )
        conn.commit()


def set_salary(user_id, salary):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("UPDATE users SET monthly_salary = ? WHERE user_id = ?", (salary, user_id))
        conn.commit()


def set_norm_days(user_id, norm_days):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("UPDATE users SET norm_days = ? WHERE user_id = ?", (norm_days, user_id))
        conn.commit()


def delete_user(user_id):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("DELETE FROM attendance WHERE user_id = ?", (user_id,))
        cursor.execute("DELETE FROM advances WHERE user_id = ?", (user_id,))
        cursor.execute("DELETE FROM users WHERE user_id = ?", (user_id,))
        conn.commit()


# ---------------- ATTENDANCE ----------------
def get_today_attendance(user_id, today_str):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT id, check_in_time, check_out_time, break_start, break_end, break_minutes "
            "FROM attendance WHERE user_id = ? AND date = ?", (user_id, today_str)
        )
        return cursor.fetchone()


def checkin(user_id, today_str, current_time_str, lateness):
    with get_db() as conn:
        cursor = conn.cursor()
        try:
            cursor.execute(
                "INSERT INTO attendance (user_id, date, check_in_time, lateness_minutes) VALUES (?, ?, ?, ?)",
                (user_id, today_str, current_time_str, lateness)
            )
        except sqlite3.IntegrityError:
            conn.rollback()
            return None
        attendance_id = cursor.lastrowid
        conn.commit()
        return attendance_id


def checkout(user_id, today_str, current_time_str):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT id, check_in_time, break_minutes, break_start, break_end "
            "FROM attendance WHERE user_id = ? AND date = ?", (user_id, today_str)
        )
        record = cursor.fetchone()
        if not record:
            return "no_checkin", None

        record_id, check_in_time, break_minutes, break_start, break_end = record
        break_minutes = break_minutes or 0
        break_autoclosed = False
        added_break_minutes = 0
        now = get_now()

        if break_start and not break_end:
            break_start_dt = datetime.strptime(f"{today_str} {break_start}", "%Y-%m-%d %H:%M:%S").replace(tzinfo=TASHKENT_TZ)
            added_break_minutes = int((now - break_start_dt).total_seconds() / 60)
            break_minutes += added_break_minutes
            break_autoclosed = True
            cursor.execute(
                "UPDATE attendance SET break_end = ?, break_minutes = ? WHERE id = ?",
                (now.strftime("%H:%M:%S"), break_minutes, record_id)
            )

        check_in_dt = datetime.strptime(f"{today_str} {check_in_time}", "%Y-%m-%d %H:%M:%S").replace(tzinfo=TASHKENT_TZ)
        total_seconds = (now - check_in_dt).total_seconds()
        break_seconds = break_minutes * 60
        net_work_hours = round(max(0, total_seconds - break_seconds) / 3600, 2)

        cursor.execute(
            "UPDATE attendance SET check_out_time = ?, work_hours = ? WHERE id = ?",
            (current_time_str, net_work_hours, record_id)
        )
        conn.commit()

        cursor.execute("SELECT monthly_salary, norm_days FROM users WHERE user_id = ?", (user_id,))
        u_data = cursor.fetchone()

    return "ok", (net_work_hours, break_autoclosed, added_break_minutes, u_data)


def break_start_action(user_id, today_str, now_str):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT id, check_in_time, break_start, check_out_time FROM attendance WHERE user_id = ? AND date = ?",
            (user_id, today_str)
        )
        record = cursor.fetchone()
        if not record:
            return "no_checkin", None
        if record[3] is not None:
            return "already_checked_out", None
        if record[2] is not None:
            return "already_on_break", record[2]
        cursor.execute("UPDATE attendance SET break_start = ? WHERE id = ?", (now_str, record[0]))
        conn.commit()
        return "ok", None


def break_end_action(user_id, today_str, now_str, now_dt):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT id, break_start, break_end, break_minutes FROM attendance WHERE user_id = ? AND date = ?",
            (user_id, today_str)
        )
        record = cursor.fetchone()
        if not record or not record[1]:
            return "not_on_break", None
        if record[2] is not None:
            return "already_returned", None
        break_start_dt = datetime.strptime(f"{today_str} {record[1]}", "%Y-%m-%d %H:%M:%S").replace(tzinfo=TASHKENT_TZ)
        this_break_minutes = int((now_dt - break_start_dt).total_seconds() / 60)
        total_break_minutes = (record[3] or 0) + this_break_minutes
        cursor.execute(
            "UPDATE attendance SET break_end = ?, break_minutes = ? WHERE id = ?",
            (now_str, total_break_minutes, record[0])
        )
        conn.commit()
        return "ok", this_break_minutes


def save_lateness_reason(attendance_id, reason):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("UPDATE attendance SET lateness_reason = ? WHERE id = ?", (reason, attendance_id))
        conn.commit()


def get_users_without_checkin_today(today_str):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT user_id, full_name FROM users WHERE is_approved = 1")
        users = cursor.fetchall()
        pending = []
        for user_id, full_name in users:
            cursor.execute("SELECT id FROM attendance WHERE user_id = ? AND date = ?", (user_id, today_str))
            if not cursor.fetchone():
                pending.append((user_id, full_name))
        return pending


# ---------------- STATS / SALARY ----------------
def get_user_stats(user_id, month_prefix):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute('''
            SELECT COUNT(id), COALESCE(SUM(lateness_minutes), 0), COALESCE(SUM(work_hours), 0.0), COALESCE(SUM(break_minutes), 0)
            FROM attendance WHERE user_id = ? AND date LIKE ?
        ''', (user_id, f"{month_prefix}%"))
        row = cursor.fetchone()
        cursor.execute("SELECT monthly_salary, norm_days FROM users WHERE user_id = ?", (user_id,))
        salary_row = cursor.fetchone()
        return row, salary_row


def get_salary_details(user_id, month_prefix):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT full_name, monthly_salary, norm_days FROM users WHERE user_id = ?", (user_id,))
        user_info = cursor.fetchone()
        if not user_info:
            return None
        cursor.execute(
            "SELECT COALESCE(SUM(work_hours), 0.0) FROM attendance WHERE user_id = ? AND date LIKE ?",
            (user_id, f"{month_prefix}%")
        )
        total_hours = cursor.fetchone()[0]
        cursor.execute(
            "SELECT COALESCE(SUM(amount), 0.0) FROM advances WHERE user_id = ? AND date LIKE ?",
            (user_id, f"{month_prefix}%")
        )
        total_advance = cursor.fetchone()[0]
        return user_info, total_hours, total_advance


def get_month_attendance_days(user_id, month_prefix):
    """Kalendar ko'rinishi uchun: shu oydagi har bir kunlik yozuv."""
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute(
            "SELECT date, check_in_time, check_out_time, lateness_minutes, work_hours "
            "FROM attendance WHERE user_id = ? AND date LIKE ? ORDER BY date",
            (user_id, f"{month_prefix}%")
        )
        return cursor.fetchall()


def add_advance(user_id, amount, date_str):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute("INSERT INTO advances (user_id, amount, date) VALUES (?, ?, ?)", (user_id, amount, date_str))
        conn.commit()


# ---------------- EXCEL ----------------
def _style_header(ws, headers):
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(ws):
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)


def generate_excel_report(month_prefix, file_path):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute('''
            SELECT u.user_id, u.full_name, u.monthly_salary, u.norm_days
            FROM users u WHERE u.is_approved = 1 ORDER BY u.full_name COLLATE NOCASE
        ''')
        users = cursor.fetchall()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Oylik Hisobot"
        headers = ["Xodim F.I.Sh", "Belgilangan Oylik", "Norma kun", "Jami Ishlangan Soat",
                   "Hisoblangan Maosh", "Berilgan Avans", "Sof Beriladigan Oylik"]
        ws.append(headers)
        _style_header(ws, headers)

        for u_id, full_name, m_salary, norm_d in users:
            cursor.execute("SELECT COALESCE(SUM(work_hours), 0.0) FROM attendance WHERE user_id = ? AND date LIKE ?",
                           (u_id, f"{month_prefix}%"))
            w_hours = cursor.fetchone()[0] or 0.0
            cursor.execute("SELECT COALESCE(SUM(amount), 0.0) FROM advances WHERE user_id = ? AND date LIKE ?",
                           (u_id, f"{month_prefix}%"))
            adv_sum = cursor.fetchone()[0] or 0.0
            hourly_rate = m_salary / (norm_d * 8) if norm_d and norm_d > 0 else 0
            earned = w_hours * hourly_rate
            net_salary = earned - adv_sum
            ws.append([full_name, m_salary, norm_d, round(w_hours, 1), round(earned, 2), adv_sum, round(net_salary, 2)])

        _autosize(ws)
        wb.save(file_path)


def generate_daily_excel_report(date_str, file_path):
    with get_db() as conn:
        cursor = conn.cursor()
        cursor.execute('''
            SELECT u.full_name, a.check_in_time, a.check_out_time, a.lateness_minutes,
                   a.lateness_reason, a.break_minutes, a.work_hours, u.monthly_salary, u.norm_days
            FROM attendance a JOIN users u ON u.user_id = a.user_id
            WHERE a.date = ? ORDER BY u.full_name COLLATE NOCASE
        ''', (date_str,))
        rows = cursor.fetchall()
        cursor.execute('''
            SELECT u.full_name FROM users u
            WHERE u.is_approved = 1 AND u.user_id NOT IN (SELECT user_id FROM attendance WHERE date = ?)
            ORDER BY u.full_name COLLATE NOCASE
        ''', (date_str,))
        absent = [r[0] for r in cursor.fetchall()]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kunlik Hisobot"
    headers = ["Xodim F.I.Sh", "Kelgan vaqti", "Ketgan vaqti", "Kechikish (daq)",
               "Kechikish sababi", "Tanaffus (daq)", "Ishlangan soat", "Kunlik topilgan pul"]
    ws.append(headers)
    _style_header(ws, headers)

    for full_name, check_in, check_out, lateness, reason, break_min, work_hours, m_salary, norm_d in rows:
        hourly_rate = (m_salary / (norm_d * 8)) if norm_d and norm_d > 0 else 0
        daily_earned = (work_hours or 0) * hourly_rate
        ws.append([full_name, check_in or "-", check_out or "Hali ketmagan", lateness or 0,
                   reason or "", break_min or 0, round(work_hours or 0, 2), round(daily_earned, 2)])

    if absent:
        absent_fill = PatternFill(start_color="FDEAEA", end_color="FDEAEA", fill_type="solid")
        for name in absent:
            row_idx = ws.max_row + 1
            ws.append([name, "KELMADI", "-", 0, "", 0, 0, 0])
            for col_num in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_num).fill = absent_fill

    _autosize(ws)
    wb.save(file_path)
